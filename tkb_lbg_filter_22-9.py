from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import tempfile
import re
import sys
import unicodedata
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DAY_NAMES = {2: "Hai", 3: "Ba", 4: "Tư", 5: "Năm", 6: "Sáu", 7: "Bảy"}
DAY_INDEX = {
    "hai": 2,
    "ba": 3,
    "tu": 4,
    "nam": 5,
    "sau": 6,
    "bay": 7,
}
SESSION_INDEX = {"Sáng": 0, "Chiều": 1}
DEFAULT_TEACHER = "Thành"
KHTN_SUBJECTS = {"KHTN-L", "KHTN-H", "KHTN-S"}
SCHOOL_CODES = {"PK": "Phước Kháng", "CBQ": "Cao Bá Quát"}


def compact_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def ascii_key(value: object) -> str:
    text = compact_text(value).lower().replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn")


def normalize_day(value: object) -> tuple[int, str] | None:
    if (
        not isinstance(value, bool)
        and isinstance(value, (int, float))
        and int(value) == value
        and 2 <= int(value) <= 7
    ):
        day_index = int(value)
        return day_index, DAY_NAMES[day_index]
    key = ascii_key(value)
    key = re.sub(r"^thu\s*", "", key).strip()
    day_index = DAY_INDEX.get(key)
    if day_index is None:
        return None
    return day_index, DAY_NAMES[day_index]


def normalize_class_name(value: object) -> str:
    text = compact_text(value)
    key = ascii_key(text)
    key = re.sub(r"^lop\s*", "", key).strip()
    key = key.replace("\\", "/")
    key = key.replace(".", "/")

    # Ưu tiên tên chữ–số; không lấy một phần của tên dài hơn.
    match = re.search(r"(?<!\w)([6-9])\s*([a-z])\s*(\d{0,2})(?!\w)", key)
    if match:
        grade, letter, number = match.groups()
        return f"{grade}{letter.upper()}{number}"

    # Dấu phân cách rõ ràng phải được xét trước dạng viết tắt 65.
    match = re.search(r"(?<!\w)(\d{1,2})\s*/\s*(\d{1,2})(?!\w)", key)
    if match:
        return f"{match.group(1)}/{match.group(2)}"

    # Giữ tương thích với cách ghi "Lớp 65" thay cho "Lớp 6/5".
    match = re.fullmatch(r"([6-9])\s*(\d{1,2})", key)
    if match:
        return f"{match.group(1)}/{match.group(2)}"
    if re.fullmatch(r"[6-9]", key):
        return key
    return ""


def most_recent_monday(today: dt.date | None = None) -> dt.date:
    """Chọn thứ Hai tuần hiện tại; từ thứ Sáu chuyển sang tuần kế tiếp."""
    today = today if today is not None else dt.date.today()
    weekday = today.weekday()
    offset = -weekday if weekday <= 3 else 7 - weekday
    return today + dt.timedelta(days=offset)


def timetable_state_path() -> Path:
    if sys.platform == "win32":
        base = Path(os.environ.get("LOCALAPPDATA") or Path.home() / "AppData" / "Local")
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support"
    else:
        base = Path(os.environ.get("XDG_CONFIG_HOME") or Path.home() / ".config")
    return base / "TKB_LBG" / "timetables.json"


def read_timetable_state(path: Path) -> list[str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return []
    if not isinstance(data, dict) or data.get("version") != 1:
        raise ValueError("Định dạng cấu hình TKB không hợp lệ.")
    paths = data.get("timetable_paths")
    if not isinstance(paths, list) or any(
        not isinstance(item, str) or not item.strip() or "\x00" in item
        or not Path(item).is_absolute() for item in paths
    ):
        raise ValueError("Danh sách đường dẫn TKB không hợp lệ.")
    return list(dict.fromkeys(paths))


def write_timetable_state(path: Path, paths: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", dir=path.parent,
            prefix="timetables-", suffix=".tmp", delete=False,
        ) as stream:
            temporary = Path(stream.name)
            json.dump({"version": 1, "timetable_paths": paths}, stream,
                      ensure_ascii=False, indent=2)
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def parse_period(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and int(value) == value:
        period = int(value)
    else:
        match = re.fullmatch(r"\s*(\d+)\s*", compact_text(value))
        if not match:
            return None
        period = int(match.group(1))
    return period if 1 <= period <= 5 else None


def is_row_visible(ws, row_number: int) -> bool:
    if hasattr(ws, "_hidden_rows"):
        return row_number not in ws._hidden_rows
    return ws.row_dimensions[row_number].hidden is not True


def is_column_visible(ws, column_number: int) -> bool:
    if hasattr(ws, "_hidden_columns"):
        return column_number not in ws._hidden_columns
    letter = get_column_letter(column_number)
    direct = ws.column_dimensions.get(letter)
    if direct is not None and direct.hidden is True:
        return False

    # Hỗ trợ cả trường hợp Excel ẩn một nhóm cột bằng một ColumnDimension.
    for dimension in ws.column_dimensions.values():
        if dimension.hidden is not True:
            continue
        min_col = dimension.min or 0
        max_col = dimension.max or min_col
        if min_col <= column_number <= max_col:
            return False
    return True


def visible_value(ws, row_number: int, column_number: int) -> object:
    if not is_row_visible(ws, row_number):
        return None
    if not is_column_visible(ws, column_number):
        return None
    return ws.cell(row=row_number, column=column_number).value


@dataclass(frozen=True)
class ClassPair:
    class_name: str
    subject_column: int
    teacher_column: int


@dataclass(frozen=True)
class ScheduleRecord:
    day_index: int
    day_name: str
    session: str
    period: int
    class_name: str
    subject: str
    teacher: str
    source_file: str
    source_row: int
    section_header_row: int
    school_code: str
    period_inferred: bool

    @property
    def sort_key(self) -> tuple:
        return (
            self.day_index,
            SESSION_INDEX[self.session],
            self.period,
            natural_key(self.class_name),
            ascii_key(self.subject),
            ascii_key(self.teacher),
            self.school_code,
        )

    @property
    def identity_key(self) -> tuple:
        return (
            self.day_index,
            self.session,
            self.period,
            ascii_key(self.class_name),
            ascii_key(self.subject),
            ascii_key(self.teacher),
            self.school_code,
        )


@dataclass
class WorkbookMetadata:
    week: str = ""
    monday: dt.date | None = None


@dataclass
class ParseResult:
    records: list[ScheduleRecord]
    teachers: list[str]
    metadata: WorkbookMetadata
    ignored_hidden_rows: int
    ignored_hidden_columns: int
    section_count: int
    duplicate_records: int
    missing_period_rows: int


@dataclass(frozen=True)
class LBGRow:
    date_text: str
    day_display: str
    period: int
    subject: str
    ppct: str
    class_name: str
    lesson_name: str
    note: str
    day_index: int
    session: str


@dataclass
class EnrichmentResult:
    rows: list[LBGRow]
    missing_ppct: int
    missing_lesson: int
    duplicate_slots: int


@dataclass(frozen=True)
class ScheduleAlert:
    alert_type: str
    record: ScheduleRecord
    details: str = ""


def natural_key(value: object) -> tuple:
    parts = re.split(r"(\d+)", ascii_key(value))
    return tuple(int(part) if part.isdigit() else part for part in parts)


class _XlsCell:
    def __init__(self, value: object) -> None:
        self.value = value


@dataclass(frozen=True)
class _XlsMergedRange:
    min_row: int
    max_row: int
    min_col: int
    max_col: int


class _XlsWorksheet:
    """Lớp tương thích tối thiểu để bộ đọc dùng được tệp Excel 97–2003."""

    def __init__(self, sheet) -> None:
        self._sheet = sheet
        self.max_row = sheet.nrows
        self.max_column = sheet.ncols
        self._hidden_rows = {
            row_index + 1
            for row_index, info in sheet.rowinfo_map.items()
            if info.hidden
        }
        self._hidden_columns = {
            column_index + 1
            for column_index, info in sheet.colinfo_map.items()
            if info.hidden
        }
        self._merged_ranges = [
            _XlsMergedRange(
                min_row=row_start + 1,
                max_row=row_end,
                min_col=column_start + 1,
                max_col=column_end,
            )
            for row_start, row_end, column_start, column_end in sheet.merged_cells
        ]

    def cell(self, row: int, column: int) -> _XlsCell:
        if row < 1 or column < 1 or row > self.max_row or column > self.max_column:
            return _XlsCell(None)
        value = self._sheet.cell_value(row - 1, column - 1)
        return _XlsCell(None if value == "" else value)


def load_schedule_sheet(path: Path):
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError(
                "Để đọc tệp .xls, hãy chạy: py -m pip install xlrd"
            ) from exc
        workbook = xlrd.open_workbook(path, formatting_info=True)
        return _XlsWorksheet(workbook.sheet_by_index(0))

    workbook = load_workbook(path, data_only=True, read_only=False)
    return workbook[workbook.sheetnames[0]]


def merged_range_starting_at(ws, row_number: int, column_number: int):
    merged_ranges = (
        ws._merged_ranges if hasattr(ws, "_merged_ranges") else ws.merged_cells.ranges
    )
    for merged_range in merged_ranges:
        if (
            merged_range.min_row <= row_number <= merged_range.max_row
            and merged_range.min_col == column_number
        ):
            return merged_range
    return None


def detect_class_pairs(ws, header_row: int) -> list[ClassPair]:
    pairs: list[ClassPair] = []
    column = 3

    while column <= ws.max_column:
        if not is_column_visible(ws, column):
            column += 1
            continue

        cell_value = visible_value(ws, header_row, column)
        class_name = normalize_class_name(cell_value)
        merged_range = merged_range_starting_at(ws, header_row, column)

        if class_name:
            teacher_column = column + 1
            if merged_range is not None and merged_range.max_col >= column + 1:
                teacher_column = column + 1

            if (
                teacher_column <= ws.max_column
                and is_column_visible(ws, teacher_column)
            ):
                pairs.append(
                    ClassPair(
                        class_name=class_name,
                        subject_column=column,
                        teacher_column=teacher_column,
                    )
                )

        if merged_range is not None:
            column = max(column + 1, merged_range.max_col + 1)
        else:
            column += 2 if class_name else 1

    return pairs


def detect_header_rows(ws) -> list[tuple[int, list[ClassPair]]]:
    headers: list[tuple[int, list[ClassPair]]] = []
    for row_number in range(1, ws.max_row + 1):
        if not is_row_visible(ws, row_number):
            continue
        col_a = ascii_key(visible_value(ws, row_number, 1))
        col_b = ascii_key(visible_value(ws, row_number, 2))
        if col_a != "thu" or "tiet" not in col_b:
            continue
        pairs = detect_class_pairs(ws, row_number)
        if pairs:
            headers.append((row_number, pairs))
    return headers


def detect_session(ws, header_row: int) -> str | None:
    # Tiêu đề buổi nằm phía trên dòng "Thứ | Tiết". Chọn nhãn gần nhất.
    for row_number in range(header_row - 1, max(0, header_row - 12), -1):
        if not is_row_visible(ws, row_number):
            continue
        for column_number in range(1, ws.max_column + 1):
            if not is_column_visible(ws, column_number):
                continue
            key = ascii_key(ws.cell(row=row_number, column=column_number).value)
            if re.search(r"\bchieu\b", key):
                return "Chiều"
            if re.search(r"\bsang\b", key):
                return "Sáng"
    return None


def extract_metadata(ws) -> WorkbookMetadata:
    fragments: list[str] = []
    for row_number in range(1, min(ws.max_row, 15) + 1):
        if not is_row_visible(ws, row_number):
            continue
        for column_number in range(1, ws.max_column + 1):
            if not is_column_visible(ws, column_number):
                continue
            value = compact_text(ws.cell(row=row_number, column=column_number).value)
            if value:
                fragments.append(value)

    text = " ".join(fragments)
    week_match = re.search(r"tuần\s*0*(\d+)", text, re.IGNORECASE)
    date_match = re.search(
        r"(?:từ\s+)?ngày\s*(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{4})",
        text,
        re.IGNORECASE,
    )

    metadata = WorkbookMetadata()
    if week_match:
        metadata.week = week_match.group(1)
    if date_match:
        try:
            metadata.monday = dt.date(
                int(date_match.group(3)),
                int(date_match.group(2)),
                int(date_match.group(1)),
            )
        except ValueError:
            metadata.monday = None
    return metadata


def detect_school_code(ws, file_name: str) -> str:
    """Nhận diện điểm trường từ tên tệp, sau đó đối chiếu nội dung Sheet 1."""
    file_key = ascii_key(Path(file_name).stem)
    if "phuoc khang" in file_key or re.search(
        r"(^|[^a-z0-9])pk([^a-z0-9]|$)", file_key
    ):
        return "PK"
    if "cao ba quat" in file_key or re.search(
        r"(^|[^a-z0-9])cbq([^a-z0-9]|$)", file_key
    ):
        return "CBQ"

    fragments: list[str] = []
    for row_number in range(1, min(ws.max_row, 20) + 1):
        for column_number in range(1, ws.max_column + 1):
            value = compact_text(ws.cell(row_number, column_number).value)
            if value:
                fragments.append(value)
    content_key = ascii_key(" ".join(fragments))
    if "phuoc khang" in content_key:
        return "PK"
    if "cao ba quat" in content_key:
        return "CBQ"
    # Dấu hiệu cấu trúc của mẫu CBQ khi tên tệp bị đổi hoặc không ghi điểm trường.
    if (
        "buoi sang" in content_key
        and "6/1" in content_key
        and "9/2" in content_key
        and ws.max_column <= 30
    ):
        return "CBQ"
    return ""


def row_has_schedule_data(ws, row_number: int, class_pairs: list[ClassPair]) -> bool:
    return any(
        compact_text(visible_value(ws, row_number, pair.subject_column)).lower()
        not in {"", "-", "x"}
        and compact_text(visible_value(ws, row_number, pair.teacher_column)).lower()
        not in {"", "-", "x"}
        for pair in class_pairs
    )


def parse_schedule_file(file_path: str | Path) -> ParseResult:
    path = Path(file_path)
    ws = load_schedule_sheet(path)
    school_code = detect_school_code(ws, path.name)

    hidden_rows = {
        row_number
        for row_number in range(1, ws.max_row + 1)
        if not is_row_visible(ws, row_number)
    }
    hidden_columns = {
        column_number
        for column_number in range(1, ws.max_column + 1)
        if not is_column_visible(ws, column_number)
    }

    headers = detect_header_rows(ws)
    if not headers:
        raise ValueError(
            f"Không tìm thấy bảng thời khóa biểu hợp lệ trong Sheet 1 của {path.name}."
        )

    records: list[ScheduleRecord] = []
    missing_period_rows = 0
    for header_index, (header_row, class_pairs) in enumerate(headers):
        session = detect_session(ws, header_row)
        if session is None:
            # Một số TKB điểm trường không ghi tên buổi: bảng đầu là Sáng,
            # bảng tiếp theo là Chiều.
            session = "Sáng" if header_index == 0 else "Chiều"

        end_row = (
            headers[header_index + 1][0] - 1
            if header_index + 1 < len(headers)
            else ws.max_row
        )
        current_day: tuple[int, str] | None = None
        last_period: int | None = None

        for row_number in range(header_row + 1, end_row + 1):
            if row_number in hidden_rows:
                continue

            raw_day = visible_value(ws, row_number, 1)
            parsed_day = normalize_day(raw_day)
            if parsed_day is not None:
                current_day = parsed_day
                last_period = None

            period = parse_period(visible_value(ws, row_number, 2))
            period_inferred = False
            if (
                period is None
                and current_day is not None
                and row_has_schedule_data(ws, row_number, class_pairs)
            ):
                missing_period_rows += 1
                if last_period is not None and last_period < 5:
                    period = last_period + 1
                    period_inferred = True

            if current_day is None or period is None:
                continue
            last_period = period

            for pair in class_pairs:
                subject = compact_text(
                    visible_value(ws, row_number, pair.subject_column)
                )
                teacher = compact_text(
                    visible_value(ws, row_number, pair.teacher_column)
                )
                if (
                    subject.lower() in {"", "-", "x"}
                    or teacher.lower() in {"", "-", "x"}
                ):
                    continue

                records.append(
                    ScheduleRecord(
                        day_index=current_day[0],
                        day_name=current_day[1],
                        session=session,
                        period=period,
                        class_name=pair.class_name,
                        subject=subject,
                        teacher=teacher,
                        source_file=path.name,
                        source_row=row_number,
                        section_header_row=header_row,
                        school_code=school_code,
                        period_inferred=period_inferred,
                    )
                )

    teachers = order_teachers(records)
    return ParseResult(
        records=records,
        teachers=teachers,
        metadata=extract_metadata(ws),
        ignored_hidden_rows=len(hidden_rows),
        ignored_hidden_columns=len(hidden_columns),
        section_count=len(headers),
        duplicate_records=0,
        missing_period_rows=missing_period_rows,
    )


def merge_parse_results(results: Iterable[ParseResult]) -> ParseResult:
    result_list = list(results)
    unique_records: list[ScheduleRecord] = []
    seen: set[tuple] = set()
    duplicate_records = sum(result.duplicate_records for result in result_list)
    missing_period_rows = sum(result.missing_period_rows for result in result_list)

    for result in result_list:
        for record in result.records:
            if record.identity_key in seen:
                duplicate_records += 1
                continue
            seen.add(record.identity_key)
            unique_records.append(record)

    unique_records.sort(key=lambda record: record.sort_key)
    teachers = order_teachers(unique_records)

    metadata = WorkbookMetadata()
    for result in result_list:
        if not metadata.week and result.metadata.week:
            metadata.week = result.metadata.week
        if metadata.monday is None and result.metadata.monday is not None:
            metadata.monday = result.metadata.monday

    return ParseResult(
        records=unique_records,
        teachers=teachers,
        metadata=metadata,
        ignored_hidden_rows=sum(r.ignored_hidden_rows for r in result_list),
        ignored_hidden_columns=sum(r.ignored_hidden_columns for r in result_list),
        section_count=sum(r.section_count for r in result_list),
        duplicate_records=duplicate_records,
        missing_period_rows=missing_period_rows,
    )


def parse_date(value: str) -> dt.date | None:
    text = compact_text(value)
    if not text:
        return None
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError("Ngày thứ Hai phải có dạng dd/mm/yyyy hoặc yyyy-mm-dd.")


def day_date(monday: dt.date | None, day_index: int) -> str:
    if monday is None:
        return ""
    return (monday + dt.timedelta(days=day_index - 2)).strftime("%d/%m/%Y")


def filter_teacher(
    records: Iterable[ScheduleRecord], teacher: str
) -> list[ScheduleRecord]:
    return filter_teachers(records, [teacher])


def filter_teachers(
    records: Iterable[ScheduleRecord], teachers: Iterable[str]
) -> list[ScheduleRecord]:
    """Gộp lịch của nhiều tên GV và bỏ dòng giống nhau do lỗi đặt tên."""
    teacher_names = [compact_text(name) for name in teachers if compact_text(name)]
    selected: list[ScheduleRecord] = []
    seen: dict[tuple, int] = {}

    # Duyệt theo thứ tự tên người dùng chọn để tên chính được ưu tiên giữ lại.
    record_list = list(records)
    for teacher in teacher_names:
        for record in record_list:
            if record.teacher != teacher:
                continue
            alias_identity = (
                record.day_index,
                record.session,
                record.period,
                ascii_key(record.class_name),
                canonical_subject(record.subject),
                record.school_code,
            )
            if alias_identity in seen:
                existing_index = seen[alias_identity]
                if selected[existing_index].period_inferred and not record.period_inferred:
                    selected[existing_index] = record
                continue
            seen[alias_identity] = len(selected)
            selected.append(record)

    selected.sort(key=lambda record: record.sort_key)
    return selected


def collect_schedule_alerts(
    records: Iterable[ScheduleRecord],
) -> list[ScheduleAlert]:
    record_list = list(records)
    slot_records: dict[tuple[int, str, int], list[ScheduleRecord]] = {}
    for record in record_list:
        key = (record.day_index, record.session, record.period)
        slot_records.setdefault(key, []).append(record)

    alerts: list[ScheduleAlert] = []
    for records_in_slot in slot_records.values():
        if len(records_in_slot) > 1:
            for record in records_in_slot:
                counterparts = [
                    other
                    for other in records_in_slot
                    if other is not record
                ]
                details = "; ".join(
                    f"Lớp {other.class_name} – {other.subject} – GV {other.teacher}"
                    for other in counterparts
                )
                alerts.append(ScheduleAlert("TRÙNG TIẾT", record, details))
    alerts.extend(
        ScheduleAlert(
            f"THIẾU TIẾT – SUY RA {record.period}",
            record,
            f"Ô số tiết bị trống; đã suy ra tiết {record.period} từ thứ tự các dòng.",
        )
        for record in record_list
        if record.period_inferred
    )
    alerts.sort(
        key=lambda alert: (
            alert.record.sort_key,
            0 if alert.alert_type == "TRÙNG TIẾT" else 1,
        )
    )
    return alerts


def collect_alias_excess_alerts(
    records: Iterable[ScheduleRecord], teachers: Iterable[str]
) -> list[ScheduleAlert]:
    teacher_set = {compact_text(name) for name in teachers if compact_text(name)}
    groups: dict[tuple, list[ScheduleRecord]] = {}
    for record in records:
        if record.teacher not in teacher_set:
            continue
        key = (
            record.day_index,
            record.session,
            record.period,
            ascii_key(record.class_name),
            canonical_subject(record.subject),
            record.school_code,
        )
        groups.setdefault(key, []).append(record)

    alerts: list[ScheduleAlert] = []
    for duplicate_records in groups.values():
        if len(duplicate_records) <= 1:
            continue
        for record in duplicate_records:
            other_names = sorted(
                {other.teacher for other in duplicate_records if other is not record},
                key=natural_key,
            )
            details = "Cùng lớp, môn và khung giờ với tên GV: " + ", ".join(
                other_names
            )
            alerts.append(
                ScheduleAlert("THỪA TIẾT (TRÙNG DÒNG)", record, details)
            )
    alerts.sort(key=lambda alert: alert.record.sort_key)
    return alerts


def collect_all_teacher_duplicate_alerts(
    records: Iterable[ScheduleRecord],
) -> list[ScheduleAlert]:
    """Chỉ coi là trùng khi đúng cùng tên GV và cùng ngày/buổi/tiết."""
    groups: dict[tuple[str, int, str, int], list[ScheduleRecord]] = {}
    for record in records:
        key = (record.teacher, record.day_index, record.session, record.period)
        groups.setdefault(key, []).append(record)

    alerts: list[ScheduleAlert] = []
    for records_in_slot in groups.values():
        if len(records_in_slot) <= 1:
            continue
        for record in records_in_slot:
            details = "; ".join(
                f"Lớp {other.class_name} – {other.subject}"
                for other in records_in_slot
                if other is not record
            )
            alerts.append(ScheduleAlert("TRÙNG TIẾT", record, details))
    alerts.sort(key=lambda alert: (natural_key(alert.record.teacher), alert.record.sort_key))
    return alerts


def safe_filename(value: str) -> str:
    value = re.sub(r"[<>:\"/\\|?*]", "_", compact_text(value))
    return value or "Giao_vien"


def normalize_number_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)) and int(value) == value:
        return str(int(value))
    return compact_text(value)


def canonical_subject(value: object) -> str:
    key = re.sub(r"\s+", " ", ascii_key(value)).strip()
    aliases = {
        "ly": "KHTN-L",
        "vat ly": "KHTN-L",
        "khtn-l": "KHTN-L",
        "khtn (l)": "KHTN-L",
        "hoa": "KHTN-H",
        "hoa hoc": "KHTN-H",
        "khtn-h": "KHTN-H",
        "khtn (h)": "KHTN-H",
        "sinh": "KHTN-S",
        "sinh hoc": "KHTN-S",
        "khtn-s": "KHTN-S",
        "khtn (s)": "KHTN-S",
        "cn": "C.NGHỆ",
        "c.nghe": "C.NGHỆ",
        "cong nghe": "C.NGHỆ",
        "hdtn": "HĐTNHN",
        "hdtnhn": "HĐTNHN",
        "hoat dong trai nghiem": "HĐTNHN",
        "hoat dong trai nghiem huong nghiep": "HĐTNHN",
    }
    return aliases.get(key, compact_text(value).upper())


def order_teachers(records: Iterable[ScheduleRecord]) -> list[str]:
    """Ưu tiên Thành, rồi giáo viên dạy KHTN, sau đó các giáo viên còn lại."""
    record_list = list(records)
    all_teachers = {record.teacher for record in record_list if record.teacher}
    science_teachers = {
        record.teacher
        for record in record_list
        if canonical_subject(record.subject) in KHTN_SUBJECTS
    }

    ordered: list[str] = []
    if DEFAULT_TEACHER in all_teachers:
        ordered.append(DEFAULT_TEACHER)

    ordered.extend(
        sorted(science_teachers - set(ordered), key=natural_key)
    )
    ordered.extend(sorted(all_teachers - set(ordered), key=natural_key))
    return ordered


def find_required_columns(
    ws, required: dict[str, tuple[str, ...]]
) -> tuple[int, dict[str, int]]:
    for row_number in range(1, min(ws.max_row, 20) + 1):
        normalized = {
            column_number: ascii_key(ws.cell(row_number, column_number).value)
            for column_number in range(1, ws.max_column + 1)
        }
        found: dict[str, int] = {}
        for field, aliases in required.items():
            for column_number, header in normalized.items():
                if any(alias in header for alias in aliases):
                    found[field] = column_number
                    break
        if len(found) == len(required):
            return row_number, found
    missing = ", ".join(required)
    raise ValueError(f"Không tìm thấy đủ các cột bắt buộc: {missing}.")


def load_ppct_map(file_path: str | Path) -> dict[tuple[str, str, str], list[str]]:
    workbook = load_workbook(file_path, data_only=True, read_only=False)
    ws = workbook[workbook.sheetnames[0]]
    header_row, columns = find_required_columns(
        ws,
        {
            "grade": ("khoi",),
            "subject": ("mon",),
            "week": ("tuan",),
            "ppct": ("tiet ppct",),
        },
    )

    result: dict[tuple[str, str, str], list[str]] = {}
    for row_number in range(header_row + 1, ws.max_row + 1):
        grade_match = re.search(
            r"\d{1,2}", compact_text(ws.cell(row_number, columns["grade"]).value)
        )
        subject = canonical_subject(ws.cell(row_number, columns["subject"]).value)
        week = normalize_number_text(ws.cell(row_number, columns["week"]).value)
        raw_ppct = normalize_number_text(ws.cell(row_number, columns["ppct"]).value)
        if not grade_match or not subject or not week or not raw_ppct:
            continue
        periods = [
            normalize_number_text(part)
            for part in re.split(r"[,;]", raw_ppct)
            if normalize_number_text(part)
        ]
        key = (grade_match.group(0), subject, week)
        result.setdefault(key, []).extend(periods)
    return result


def load_lesson_map(file_path: str | Path) -> dict[tuple[str, str, str], str]:
    workbook = load_workbook(file_path, data_only=True, read_only=False)
    ws = workbook[workbook.sheetnames[0]]
    header_row, columns = find_required_columns(
        ws,
        {
            "grade": ("khoi",),
            "subject": ("mon",),
            "ppct": ("tiet ppct",),
            "lesson": ("ten bai", "noi dung day"),
        },
    )

    result: dict[tuple[str, str, str], str] = {}
    for row_number in range(header_row + 1, ws.max_row + 1):
        grade_match = re.search(
            r"\d{1,2}", compact_text(ws.cell(row_number, columns["grade"]).value)
        )
        subject = canonical_subject(ws.cell(row_number, columns["subject"]).value)
        ppct = normalize_number_text(ws.cell(row_number, columns["ppct"]).value)
        lesson = compact_text(ws.cell(row_number, columns["lesson"]).value)
        if not grade_match or not subject or not ppct or not lesson:
            continue
        result[(grade_match.group(0), subject, ppct)] = lesson
    return result


def build_lbg_rows(
    records: list[ScheduleRecord],
    week: str,
    monday: dt.date | None,
    ppct_file: str | Path,
    lesson_file: str | Path,
) -> EnrichmentResult:
    ppct_map = load_ppct_map(ppct_file)
    lesson_map = load_lesson_map(lesson_file)
    used: dict[tuple[str, str, str, str], int] = {}
    rows: list[LBGRow] = []
    missing_ppct = 0
    missing_lesson = 0
    slot_counts: dict[tuple[int, str, int], int] = {}
    for record in records:
        slot_key = (record.day_index, record.session, record.period)
        slot_counts[slot_key] = slot_counts.get(slot_key, 0) + 1
    duplicate_slots = sum(count > 1 for count in slot_counts.values())

    for record in sorted(records, key=lambda item: item.sort_key):
        grade_digits = re.match(r"\d{1,2}", record.class_name)
        grade = grade_digits.group(0) if grade_digits else record.class_name
        subject_key = canonical_subject(record.subject)
        ppct_key = (grade, subject_key, normalize_number_text(week))
        usage_key = (*ppct_key, record.class_name)
        index = used.get(usage_key, 0)
        ppct_values = ppct_map.get(ppct_key, [])
        ppct = ppct_values[index] if index < len(ppct_values) else ""
        if ppct:
            used[usage_key] = index + 1
        else:
            missing_ppct += 1

        lesson_name = lesson_map.get((grade, subject_key, ppct), "") if ppct else ""
        if ppct and not lesson_name:
            missing_lesson += 1

        warnings: list[str] = [record.school_code] if record.school_code else []
        if slot_counts[(record.day_index, record.session, record.period)] > 1:
            warnings.append("TRÙNG TIẾT")
        if record.period_inferred:
            warnings.append(f"THIẾU TIẾT (SUY RA {record.period})")
        if ppct and not lesson_name:
            warnings.append("THIẾU TÊN BÀI")

        session_prefix = "S" if record.session == "Sáng" else "C"
        rows.append(
            LBGRow(
                date_text=day_date(monday, record.day_index),
                day_display=f"{session_prefix}.{record.day_name}",
                period=record.period,
                subject=record.subject,
                ppct=ppct,
                class_name=record.class_name,
                lesson_name=lesson_name,
                note="; ".join(warnings),
                day_index=record.day_index,
                session=record.session,
            )
        )

    return EnrichmentResult(rows, missing_ppct, missing_lesson, duplicate_slots)


def _shift_template_rows(ws, insert_at: int, amount: int) -> None:
    if amount <= 0:
        return

    merged_to_shift = []
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= insert_at:
            merged_to_shift.append(
                (
                    merged_range.min_row,
                    merged_range.max_row,
                    merged_range.min_col,
                    merged_range.max_col,
                )
            )
            ws.unmerge_cells(str(merged_range))

    row_dimensions = {}
    for row_number in list(ws.row_dimensions):
        if row_number >= insert_at:
            row_dimensions[row_number] = copy(ws.row_dimensions[row_number])
            del ws.row_dimensions[row_number]

    ws.insert_rows(insert_at, amount=amount)

    for old_row, dimension in row_dimensions.items():
        new_row = old_row + amount
        dimension.index = new_row
        ws.row_dimensions[new_row] = dimension

    for min_row, max_row, min_col, max_col in merged_to_shift:
        ws.merge_cells(
            start_row=min_row + amount,
            end_row=max_row + amount,
            start_column=min_col,
            end_column=max_col,
        )


def export_lbg(
    rows: list[LBGRow],
    output_path: str | Path,
    template_file: str | Path,
    week: str,
    monday: dt.date | None,
) -> None:
    if not rows:
        raise ValueError("Không có tiết dạy để xuất.")

    workbook = load_workbook(template_file)
    ws = workbook[workbook.sheetnames[0]]

    sunday = monday + dt.timedelta(days=6) if monday else None
    if week and monday and sunday:
        title = (
            f"LBG Tuần {week}, Từ ngày {monday.strftime('%d-%m-%Y')} "
            f"đến ngày {sunday.strftime('%d-%m-%Y')}"
        )
    elif week:
        title = f"LBG Tuần {week}"
    else:
        title = "LBG Hoàn chỉnh"
    ws["A1"] = title

    section_two_row = None
    for row_number in range(1, ws.max_row + 1):
        value = compact_text(ws.cell(row=row_number, column=1).value).upper()
        if value.startswith("II."):
            section_two_row = row_number
            break
    if section_two_row is None:
        raise ValueError("Tệp A4 không có mục II để xác định cuối bảng giảng dạy.")

    template_data_row = 5
    data_end_row = section_two_row - 2
    capacity = data_end_row - template_data_row + 1
    if len(rows) > capacity:
        _shift_template_rows(ws, section_two_row - 1, len(rows) - capacity)

    day_colors = {
        2: "FFC0D9EF",
        3: "FFC6EFCE",
        4: "FFFFFFCC",
        5: "FFFFE699",
        6: "FFD9D2E9",
        7: "FFE0E0E0",
    }
    template_styles = [
        copy(ws.cell(row=template_data_row, column=column_number)._style)
        for column_number in range(1, 9)
    ]
    template_height = ws.row_dimensions[template_data_row].height
    # Mẫu gốc có cột A hơi hẹp so với ngày dạng dd/mm/yyyy.
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 11.5)

    for output_row, row_data in enumerate(rows, start=template_data_row):
        values = [
            row_data.date_text,
            row_data.day_display,
            row_data.period,
            row_data.subject,
            row_data.ppct,
            row_data.class_name,
            row_data.lesson_name,
            row_data.note,
        ]
        fill = PatternFill(
            start_color=day_colors[row_data.day_index],
            end_color=day_colors[row_data.day_index],
            fill_type="solid",
        )
        for column_number, value in enumerate(values, start=1):
            cell = ws.cell(row=output_row, column=column_number, value=value)
            cell._style = copy(template_styles[column_number - 1])
            cell.fill = fill
            if column_number in (7, 8):
                wrapped_alignment = copy(cell.alignment)
                wrapped_alignment.wrap_text = True
                cell.alignment = wrapped_alignment
        base_height = template_height or 15
        lesson_lines = max(1, (len(row_data.lesson_name) + 59) // 60)
        note_lines = max(1, (len(row_data.note) + 11) // 12)
        content_lines = max(lesson_lines, note_lines)
        ws.row_dimensions[output_row].height = max(
            base_height, min(content_lines, 4) * 15
        )

    workbook.save(output_path)


def export_all_teacher_conflicts(
    alerts: list[ScheduleAlert],
    output_path: str | Path,
    monday: dt.date | None,
) -> None:
    if not alerts:
        raise ValueError("Không có tiết trùng để xuất.")

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Tiết trùng"
    ws.sheet_view.showGridLines = False

    headers = [
        "STT",
        "Giáo viên",
        "Ngày",
        "Thứ",
        "Buổi",
        "Tiết",
        "Lớp",
        "Môn",
        "Điểm trường",
        "Trùng với lớp/môn",
        "Tệp nguồn",
        "Dòng nguồn",
    ]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(1, 1, "DANH SÁCH TẤT CẢ TIẾT TRÙNG CỦA GIÁO VIÊN")
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="9C0006")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    header_fill = PatternFill("solid", fgColor="C00000")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column_number, header in enumerate(headers, start=1):
        cell = ws.cell(2, column_number, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[2].height = 32

    row_fill = PatternFill("solid", fgColor="FFC7CE")
    for index, alert in enumerate(alerts, start=1):
        record = alert.record
        values = [
            index,
            record.teacher,
            day_date(monday, record.day_index),
            record.day_name,
            record.session,
            record.period,
            record.class_name,
            record.subject,
            record.school_code or "Trường chính",
            alert.details,
            record.source_file,
            record.source_row,
        ]
        output_row = index + 2
        for column_number, value in enumerate(values, start=1):
            cell = ws.cell(output_row, column_number, value)
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column_number in (2, 8, 10, 11) else "center",
                vertical="center",
                wrap_text=True,
            )
        ws.row_dimensions[output_row].height = 30

    widths = [7, 16, 13, 9, 10, 8, 9, 18, 14, 36, 38, 12]
    for column_number, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column_number)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:L{len(alerts) + 2}"
    ws.print_title_rows = "1:2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(output_path)


class TimetableFilterApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Lọc thời khóa biểu và xuất lịch báo giảng")
        self.geometry("1180x820")
        self.minsize(980, 700)

        self.timetable_paths: list[str] = []
        self.ppct_path = tk.StringVar()
        self.lesson_path = tk.StringVar()
        self.template_path = tk.StringVar()
        self.teacher_var = tk.StringVar()
        self.selected_teachers: list[str] = []
        self.teacher_selection_var = tk.StringVar(value="Tên GV đang gộp: chưa chọn")
        self.week_var = tk.StringVar()
        self.monday_var = tk.StringVar(value=most_recent_monday().strftime("%d/%m/%Y"))
        self.status_var = tk.StringVar(value="Sẵn sàng. Hãy thêm các tệp thời khóa biểu.")

        self.merged_result: ParseResult | None = None
        self.filtered_records: list[ScheduleRecord] = []
        self.enrichment_result: EnrichmentResult | None = None
        self.alert_window: tk.Toplevel | None = None

        self._build_ui()
        self._discover_support_files()
        self._restore_timetable_files()

    def _restore_timetable_files(self) -> None:
        try:
            self.timetable_paths = read_timetable_state(timetable_state_path())
        except (OSError, ValueError) as exc:
            self.status_var.set(f"Không đọc được cấu hình TKB: {exc}. Hãy thêm lại tệp.")
            return
        if self.timetable_paths:
            self._refresh_timetable_list()
            missing = sum(not Path(path).is_file() for path in self.timetable_paths)
            self.status_var.set(
                f"Đã khôi phục {len(self.timetable_paths)} tệp TKB. "
                + (f"{missing} tệp không còn ở vị trí cũ; hãy xóa mục cũ và thêm lại. " if missing else "")
                + "Nhấn Đọc và tổng hợp TKB để sử dụng."
            )

    def _save_timetable_files(self) -> None:
        try:
            write_timetable_state(timetable_state_path(), self.timetable_paths)
        except (OSError, ValueError) as exc:
            messagebox.showwarning(
                "Không lưu được danh sách TKB",
                f"Các tệp vẫn dùng được trong phiên này, nhưng chưa lưu cho lần sau.\n{exc}",
            )

    def _build_ui(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        input_frame = ttk.LabelFrame(self, text="1. Chọn dữ liệu", padding=12)
        input_frame.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 6))
        input_frame.columnconfigure(1, weight=1)

        ttk.Label(input_frame, text="Danh sách TKB:").grid(
            row=0, column=0, sticky="nw", pady=4
        )
        timetable_frame = ttk.Frame(input_frame)
        timetable_frame.grid(row=0, column=1, sticky="ew", padx=8, pady=4)
        timetable_frame.columnconfigure(0, weight=1)
        self.timetable_list = tk.Listbox(
            timetable_frame,
            height=4,
            selectmode=tk.EXTENDED,
            exportselection=False,
        )
        timetable_scroll = ttk.Scrollbar(
            timetable_frame, orient="vertical", command=self.timetable_list.yview
        )
        self.timetable_list.configure(yscrollcommand=timetable_scroll.set)
        self.timetable_list.grid(row=0, column=0, sticky="ew")
        timetable_scroll.grid(row=0, column=1, sticky="ns")

        timetable_buttons = ttk.Frame(input_frame)
        timetable_buttons.grid(row=0, column=2, sticky="n", pady=4)
        ttk.Button(
            timetable_buttons,
            text="Thêm TKB...",
            command=self._add_timetable_files,
        ).pack(fill="x")
        ttk.Button(
            timetable_buttons,
            text="Xóa tệp chọn",
            command=self._remove_timetable_files,
        ).pack(fill="x", pady=(4, 0))
        ttk.Button(
            timetable_buttons,
            text="Xóa tất cả",
            command=self._clear_timetable_files,
        ).pack(fill="x", pady=(4, 0))

        support_files = [
            ("A2 – Tuần/Tiết PPCT:", self.ppct_path),
            ("A3 – Tiết PPCT/Bài học:", self.lesson_path),
            ("A4 – Mẫu LBG:", self.template_path),
        ]
        for row_number, (label, variable) in enumerate(support_files, start=1):
            ttk.Label(input_frame, text=label).grid(
                row=row_number, column=0, sticky="w", pady=3
            )
            ttk.Entry(input_frame, textvariable=variable, state="readonly").grid(
                row=row_number, column=1, sticky="ew", padx=8, pady=3
            )
            ttk.Button(
                input_frame,
                text="Chọn tệp...",
                command=lambda current=variable: self._choose_file(current),
            ).grid(row=row_number, column=2, pady=3)

        controls = ttk.Frame(input_frame)
        controls.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(10, 0))
        ttk.Button(controls, text="Đọc và tổng hợp TKB", command=self._load_files).pack(
            side="left", padx=(0, 14)
        )

        ttk.Label(controls, text="Tuần:").pack(side="left")
        ttk.Entry(controls, textvariable=self.week_var, width=7).pack(side="left", padx=(4, 12))
        ttk.Label(controls, text="Ngày thứ Hai:").pack(side="left")
        ttk.Entry(controls, textvariable=self.monday_var, width=13).pack(side="left", padx=(4, 12))

        ttk.Label(controls, text="Giáo viên:").pack(side="left")
        self.teacher_combo = ttk.Combobox(
            controls,
            textvariable=self.teacher_var,
            state="readonly",
            width=22,
        )
        self.teacher_combo.pack(side="left", padx=(4, 12))
        self.teacher_combo.bind(
            "<<ComboboxSelected>>", self._on_primary_teacher_selected
        )
        self.merge_teachers_button = ttk.Button(
            controls,
            text="Gộp tên GV...",
            command=self._choose_teacher_aliases,
            state="disabled",
        )
        self.merge_teachers_button.pack(side="left", padx=(0, 12))
        self.filter_button = ttk.Button(
            controls, text="Lọc & xem trước", command=self._filter_preview, state="disabled"
        )
        self.filter_button.pack(side="left")

        ttk.Label(
            input_frame,
            textvariable=self.teacher_selection_var,
            foreground="#1F4E78",
        ).grid(row=5, column=0, columnspan=3, sticky="w", pady=(6, 0))

        preview_frame = ttk.LabelFrame(self, text="2. Lịch dạy đã sắp xếp", padding=8)
        preview_frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=6)
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(0, weight=1)

        columns = ("date", "day", "period", "subject", "ppct", "class", "lesson", "note")
        self.tree = ttk.Treeview(preview_frame, columns=columns, show="headings")
        headings = {
            "date": "Ngày",
            "day": "Thứ",
            "period": "Tiết",
            "subject": "Môn (phân môn)",
            "ppct": "Tiết PPCT",
            "class": "Lớp",
            "lesson": "Tên bài hoặc nội dung dạy",
            "note": "Ghi chú",
        }
        widths = {
            "date": 95,
            "day": 80,
            "period": 55,
            "subject": 130,
            "ppct": 80,
            "class": 65,
            "lesson": 390,
            "note": 220,
        }
        for column in columns:
            self.tree.heading(column, text=headings[column])
            self.tree.column(column, width=widths[column], anchor="center")
        self.tree.column("subject", anchor="w")
        self.tree.column("lesson", anchor="w")
        self.tree.column("note", anchor="w")

        vertical_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=self.tree.yview)
        horizontal_scroll = ttk.Scrollbar(preview_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(
            yscrollcommand=vertical_scroll.set,
            xscrollcommand=horizontal_scroll.set,
        )
        self.tree.grid(row=0, column=0, sticky="nsew")
        vertical_scroll.grid(row=0, column=1, sticky="ns")
        horizontal_scroll.grid(row=1, column=0, sticky="ew")

        footer = ttk.Frame(self, padding=(12, 6, 12, 12))
        footer.grid(row=2, column=0, sticky="ew")
        footer.columnconfigure(0, weight=1)
        ttk.Label(footer, textvariable=self.status_var).grid(row=0, column=0, sticky="w")
        self.export_button = ttk.Button(
            footer,
            text="3. Xuất lịch báo giảng Excel",
            command=self._export,
            state="disabled",
        )
        self.all_conflicts_button = ttk.Button(
            footer,
            text="Xuất tiết trùng tất cả GV",
            command=self._export_all_teacher_conflicts,
            state="disabled",
        )
        self.all_conflicts_button.grid(row=0, column=1, sticky="e", padx=(8, 12))
        self.export_button.grid(row=0, column=2, sticky="e")

    def _choose_file(self, variable: tk.StringVar) -> None:
        selected = filedialog.askopenfilename(
            title="Chọn tệp Excel",
            filetypes=[("Excel", "*.xlsx")],
        )
        if selected:
            variable.set(selected)
            self.export_button.configure(state="disabled")
            if (
                self.merged_result is not None
                and self.ppct_path.get()
                and self.lesson_path.get()
            ):
                self._filter_preview(show_dialog=False)

    def _invalidate_timetable_result(self) -> None:
        self.merged_result = None
        self.filtered_records = []
        self.enrichment_result = None
        self.selected_teachers = []
        self.teacher_var.set("")
        self.teacher_selection_var.set("Tên GV đang gộp: chưa chọn")
        self.filter_button.configure(state="disabled")
        self.merge_teachers_button.configure(state="disabled")
        self.all_conflicts_button.configure(state="disabled")
        self.export_button.configure(state="disabled")
        self._clear_tree()
        if self.alert_window is not None and self.alert_window.winfo_exists():
            self.alert_window.destroy()
        self.alert_window = None

    def _update_teacher_selection_label(self) -> None:
        if not self.selected_teachers:
            self.teacher_selection_var.set("Tên GV đang gộp: chưa chọn")
            return
        self.teacher_selection_var.set(
            "Tên GV đang gộp: " + " + ".join(self.selected_teachers)
        )

    def _on_primary_teacher_selected(self, _event=None) -> None:
        teacher = compact_text(self.teacher_var.get())
        self.selected_teachers = [teacher] if teacher else []
        self._update_teacher_selection_label()
        self._filter_preview()

    def _choose_teacher_aliases(self) -> None:
        if self.merged_result is None or not self.merged_result.teachers:
            return

        dialog = tk.Toplevel(self)
        dialog.title("Gộp các tên của cùng một giáo viên")
        dialog.geometry("430x520")
        dialog.transient(self)
        dialog.grab_set()
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(1, weight=1)

        ttk.Label(
            dialog,
            text=(
                "Tích vào các ô của những tên được dùng cho cùng một giáo viên.\n"
                "Nhấn Áp dụng để gộp và chạy lại lịch."
            ),
            padding=10,
        ).grid(row=0, column=0, sticky="ew")

        list_frame = ttk.Frame(dialog, padding=(10, 0, 10, 8))
        list_frame.grid(row=1, column=0, sticky="nsew")
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        teacher_canvas = tk.Canvas(list_frame, highlightthickness=1)
        teacher_scroll = ttk.Scrollbar(
            list_frame, orient="vertical", command=teacher_canvas.yview
        )
        teacher_canvas.configure(yscrollcommand=teacher_scroll.set)
        teacher_canvas.grid(row=0, column=0, sticky="nsew")
        teacher_scroll.grid(row=0, column=1, sticky="ns")

        checkbox_frame = ttk.Frame(teacher_canvas, padding=6)
        checkbox_window = teacher_canvas.create_window(
            (0, 0), window=checkbox_frame, anchor="nw"
        )

        def update_scroll_region(_event=None) -> None:
            teacher_canvas.configure(scrollregion=teacher_canvas.bbox("all"))

        def fit_checkbox_width(event) -> None:
            teacher_canvas.itemconfigure(checkbox_window, width=event.width)

        checkbox_frame.bind("<Configure>", update_scroll_region)
        teacher_canvas.bind("<Configure>", fit_checkbox_width)

        def scroll_teachers(event) -> None:
            teacher_canvas.yview_scroll(int(-event.delta / 120), "units")

        teacher_canvas.bind(
            "<Enter>", lambda _event: teacher_canvas.bind_all("<MouseWheel>", scroll_teachers)
        )
        teacher_canvas.bind(
            "<Leave>", lambda _event: teacher_canvas.unbind_all("<MouseWheel>")
        )

        selected_set = set(self.selected_teachers)
        checkbox_vars: dict[str, tk.BooleanVar] = {}
        for index, teacher in enumerate(self.merged_result.teachers):
            variable = tk.BooleanVar(value=teacher in selected_set)
            checkbox_vars[teacher] = variable
            ttk.Checkbutton(
                checkbox_frame,
                text=teacher,
                variable=variable,
            ).grid(row=index, column=0, sticky="w", padx=4, pady=2)

        def close_dialog() -> None:
            teacher_canvas.unbind_all("<MouseWheel>")
            dialog.destroy()

        def clear_all_checkboxes() -> None:
            for variable in checkbox_vars.values():
                variable.set(False)

        def apply_selection() -> None:
            chosen_teachers = [
                teacher
                for teacher in self.merged_result.teachers
                if checkbox_vars[teacher].get()
            ]
            if not chosen_teachers:
                messagebox.showwarning(
                    "Chưa chọn tên",
                    "Hãy chọn ít nhất một tên giáo viên.",
                    parent=dialog,
                )
                return
            primary_teacher = compact_text(self.teacher_var.get())
            if primary_teacher in chosen_teachers:
                chosen_teachers.remove(primary_teacher)
                chosen_teachers.insert(0, primary_teacher)
            self.selected_teachers = chosen_teachers
            self.teacher_var.set(self.selected_teachers[0])
            self._update_teacher_selection_label()
            close_dialog()
            self._filter_preview()

        button_frame = ttk.Frame(dialog, padding=10)
        button_frame.grid(row=2, column=0, sticky="e")
        ttk.Button(
            button_frame,
            text="Bỏ chọn tất cả",
            command=clear_all_checkboxes,
        ).pack(side="left", padx=(0, 16))
        ttk.Button(button_frame, text="Hủy", command=close_dialog).pack(
            side="right", padx=(8, 0)
        )
        ttk.Button(button_frame, text="Áp dụng", command=apply_selection).pack(
            side="right"
        )
        dialog.protocol("WM_DELETE_WINDOW", close_dialog)

    def _refresh_timetable_list(self) -> None:
        self.timetable_list.delete(0, tk.END)
        for index, path in enumerate(self.timetable_paths, start=1):
            self.timetable_list.insert(tk.END, f"{index}. {Path(path).name}" + (" [Không tìm thấy]" if not Path(path).is_file() else ""))
        self.status_var.set(f"Đã thêm {len(self.timetable_paths)} tệp TKB.")

    def _add_timetable_files(self) -> None:
        selected = filedialog.askopenfilenames(
            title="Thêm một hoặc nhiều tệp thời khóa biểu",
            filetypes=[
                ("Tệp Excel", ("*.xlsx", "*.xls")),
                ("Excel mới", "*.xlsx"),
                ("Excel 97–2003", "*.xls"),
            ],
        )
        if not selected:
            return
        known = {str(Path(path).resolve()).lower() for path in self.timetable_paths}
        for path in selected:
            normalized = str(Path(path).resolve())
            if normalized.lower() not in known:
                self.timetable_paths.append(normalized)
                known.add(normalized.lower())
        self._invalidate_timetable_result()
        self._refresh_timetable_list()
        self._save_timetable_files()

    def _remove_timetable_files(self) -> None:
        selected_indexes = list(self.timetable_list.curselection())
        if not selected_indexes:
            return
        for index in reversed(selected_indexes):
            del self.timetable_paths[index]
        self._invalidate_timetable_result()
        self._refresh_timetable_list()
        self._save_timetable_files()

    def _clear_timetable_files(self) -> None:
        self.timetable_paths.clear()
        self._invalidate_timetable_result()
        self._refresh_timetable_list()
        self._save_timetable_files()

    def _discover_support_files(self) -> None:
        folders = [Path(sys.argv[0]).resolve().parent, Path.cwd()]
        rules = [
            ("A2.", self.ppct_path),
            ("A3.", self.lesson_path),
            ("A4.", self.template_path),
        ]
        for prefix, variable in rules:
            for folder in folders:
                matches = sorted(
                    path
                    for path in folder.glob("*.xlsx")
                    if path.name.upper().startswith(prefix)
                )
                if matches:
                    variable.set(str(matches[0]))
                    break

    def _load_files(self) -> None:
        if not self.timetable_paths:
            messagebox.showerror("Thiếu tệp", "Vui lòng thêm ít nhất một tệp thời khóa biểu.")
            return

        try:
            self.status_var.set(
                f"Đang đọc {len(self.timetable_paths)} tệp và bỏ qua dữ liệu ẩn trong TKB..."
            )
            self.update_idletasks()
            available = [path for path in self.timetable_paths if Path(path).is_file()]
            missing = [path for path in self.timetable_paths if not Path(path).is_file()]
            if missing:
                messagebox.showwarning(
                    "Không tìm thấy tệp TKB",
                    "Các tệp sau đã bị di chuyển hoặc xóa; hãy xóa mục cũ và thêm lại:\n"
                    + "\n".join(missing),
                )
            if not available:
                self._invalidate_timetable_result()
                self.status_var.set("Không có tệp TKB khả dụng. Hãy thêm lại tệp.")
                return
            parse_results = [parse_schedule_file(path) for path in available]
            self.merged_result = merge_parse_results(parse_results)
            self.filtered_records = []
            self.enrichment_result = None

            self.teacher_combo.configure(values=self.merged_result.teachers)
            if self.merged_result.teachers:
                self.teacher_var.set(self.merged_result.teachers[0])
                self.selected_teachers = [self.merged_result.teachers[0]]
                self._update_teacher_selection_label()
            if self.merged_result.metadata.week:
                self.week_var.set(self.merged_result.metadata.week)
            # Giữ ngày mặc định theo thời điểm mở ứng dụng hoặc ngày người dùng nhập.
            # Ngày trong TKB có thể thuộc tuần cũ, không tự ghi đè ô ngày.

            self.filter_button.configure(state="normal")
            self.merge_teachers_button.configure(state="normal")
            self.all_conflicts_button.configure(state="normal")
            self.export_button.configure(state="disabled")
            self._clear_tree()
            self.status_var.set(
                f"Đã tổng hợp {len(self.merged_result.records)} tiết; "
                f"{len(self.merged_result.teachers)} tên giáo viên; "
                f"bỏ qua {self.merged_result.ignored_hidden_rows} hàng ẩn; "
                f"loại {self.merged_result.duplicate_records} bản ghi trùng; "
                f"thiếu số tiết: {self.merged_result.missing_period_rows}."
            )
            source_warnings: list[str] = []
            if self.merged_result.duplicate_records:
                source_warnings.append(
                    f"Đã loại {self.merged_result.duplicate_records} bản ghi trùng hoàn toàn."
                )
            if source_warnings:
                messagebox.showwarning(
                    "Cảnh báo dữ liệu TKB",
                    "\n".join(source_warnings),
                )
            if self.merged_result.teachers and self.ppct_path.get() and self.lesson_path.get():
                self._filter_preview(show_dialog=True)
        except Exception as exc:
            self.merged_result = None
            self.filter_button.configure(state="disabled")
            self.merge_teachers_button.configure(state="disabled")
            self.all_conflicts_button.configure(state="disabled")
            self.export_button.configure(state="disabled")
            self.status_var.set("Có lỗi khi đọc thời khóa biểu.")
            messagebox.showerror("Không thể đọc tệp", str(exc))

    def _clear_tree(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)

    def _show_schedule_alert_table(
        self, alerts: list[ScheduleAlert], teacher_label: str
    ) -> None:
        if self.alert_window is not None and self.alert_window.winfo_exists():
            self.alert_window.destroy()
        self.alert_window = None
        if not alerts:
            return

        dialog = tk.Toplevel(self)
        self.alert_window = dialog
        dialog.title(f"Cảnh báo lịch dạy – {teacher_label}")
        dialog.geometry("1180x500")
        dialog.minsize(900, 360)
        dialog.transient(self)
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(1, weight=1)

        duplicate_rows = sum(alert.alert_type == "TRÙNG TIẾT" for alert in alerts)
        duplicate_slots = len(
            {
                (
                    alert.record.day_index,
                    alert.record.session,
                    alert.record.period,
                )
                for alert in alerts
                if alert.alert_type == "TRÙNG TIẾT"
            }
        )
        missing_alerts = [
            alert for alert in alerts if alert.alert_type.startswith("THIẾU TIẾT")
        ]
        excess_alerts = [
            alert for alert in alerts if alert.alert_type.startswith("THỪA TIẾT")
        ]
        missing_source_rows = len(
            {
                (alert.record.source_file, alert.record.source_row)
                for alert in missing_alerts
            }
        )
        ttk.Label(
            dialog,
            text=(
                f"Giáo viên: {teacher_label}   |   "
                f"Trùng: {duplicate_slots} khung tiết ({duplicate_rows} dòng)   |   "
                f"Thừa: {len(excess_alerts)} dòng   |   "
                f"Thiếu số tiết: {missing_source_rows} dòng nguồn "
                f"({len(missing_alerts)} lịch bị ảnh hưởng)\n"
                "Hãy đối chiếu cột Tệp nguồn và Dòng nguồn để sửa thời khóa biểu gốc."
            ),
            padding=10,
            foreground="#9C0006",
        ).grid(row=0, column=0, sticky="ew")

        table_frame = ttk.Frame(dialog, padding=(10, 0, 10, 8))
        table_frame.grid(row=1, column=0, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)

        columns = (
            "alert",
            "day",
            "session",
            "period",
            "class",
            "subject",
            "teacher",
            "school",
            "details",
            "source",
            "row",
        )
        alert_tree = ttk.Treeview(
            table_frame, columns=columns, show="headings", selectmode="browse"
        )
        headings = {
            "alert": "Loại cảnh báo",
            "day": "Thứ",
            "session": "Buổi",
            "period": "Tiết",
            "class": "Lớp",
            "subject": "Môn",
            "teacher": "Tên GV trong TKB",
            "school": "Điểm trường",
            "details": "Trùng/đối chiếu với",
            "source": "Tệp nguồn",
            "row": "Dòng nguồn",
        }
        widths = {
            "alert": 170,
            "day": 60,
            "session": 70,
            "period": 50,
            "class": 60,
            "subject": 120,
            "teacher": 130,
            "school": 90,
            "details": 300,
            "source": 300,
            "row": 85,
        }
        for column in columns:
            alert_tree.heading(column, text=headings[column])
            alert_tree.column(column, width=widths[column], anchor="center")
        alert_tree.column("source", anchor="w")
        alert_tree.column("subject", anchor="w")
        alert_tree.column("details", anchor="w")
        alert_tree.tag_configure("duplicate", background="#FFC7CE")
        alert_tree.tag_configure("missing", background="#FFEB9C")
        alert_tree.tag_configure("excess", background="#FCE4D6")

        for alert in alerts:
            record = alert.record
            if alert.alert_type == "TRÙNG TIẾT":
                tag = "duplicate"
            elif alert.alert_type.startswith("THỪA TIẾT"):
                tag = "excess"
            else:
                tag = "missing"
            alert_tree.insert(
                "",
                "end",
                values=(
                    alert.alert_type,
                    record.day_name,
                    record.session,
                    record.period,
                    record.class_name,
                    record.subject,
                    record.teacher,
                    record.school_code or "Trường chính",
                    alert.details,
                    record.source_file,
                    record.source_row,
                ),
                tags=(tag,),
            )

        vertical_scroll = ttk.Scrollbar(
            table_frame, orient="vertical", command=alert_tree.yview
        )
        horizontal_scroll = ttk.Scrollbar(
            table_frame, orient="horizontal", command=alert_tree.xview
        )
        alert_tree.configure(
            yscrollcommand=vertical_scroll.set,
            xscrollcommand=horizontal_scroll.set,
        )
        alert_tree.grid(row=0, column=0, sticky="nsew")
        vertical_scroll.grid(row=0, column=1, sticky="ns")
        horizontal_scroll.grid(row=1, column=0, sticky="ew")

        def close_alerts() -> None:
            dialog.destroy()
            self.alert_window = None

        ttk.Button(dialog, text="Đã xem – Đóng", command=close_alerts).grid(
            row=2, column=0, sticky="e", padx=10, pady=(0, 10)
        )
        dialog.protocol("WM_DELETE_WINDOW", close_alerts)
        dialog.lift()
        dialog.focus_force()

    def _filter_preview(self, show_dialog: bool = True) -> None:
        if self.merged_result is None:
            return
        teachers = self.selected_teachers or [compact_text(self.teacher_var.get())]
        teachers = [teacher for teacher in teachers if teacher]
        if not teachers:
            messagebox.showerror("Chưa chọn giáo viên", "Vui lòng chọn một giáo viên.")
            return
        teacher_label = " + ".join(teachers)
        if not self.ppct_path.get() or not self.lesson_path.get():
            messagebox.showerror(
                "Thiếu tệp A2/A3",
                "Vui lòng chọn tệp A2 (PPCT) và A3 (tên bài).",
            )
            return

        try:
            monday = parse_date(self.monday_var.get())
            week = compact_text(self.week_var.get())
            teacher_set = set(teachers)
            raw_alias_count = sum(
                record.teacher in teacher_set
                for record in self.merged_result.records
            )
            self.filtered_records = filter_teachers(self.merged_result.records, teachers)
            alias_duplicates_removed = raw_alias_count - len(self.filtered_records)
            self.enrichment_result = build_lbg_rows(
                records=self.filtered_records,
                week=week,
                monday=monday,
                ppct_file=self.ppct_path.get(),
                lesson_file=self.lesson_path.get(),
            )
            self._clear_tree()
            for row_data in self.enrichment_result.rows:
                self.tree.insert(
                    "",
                    "end",
                    values=(
                        row_data.date_text,
                        row_data.day_display,
                        row_data.period,
                        row_data.subject,
                        row_data.ppct,
                        row_data.class_name,
                        row_data.lesson_name,
                        row_data.note,
                    ),
                )

            if not self.filtered_records:
                self.export_button.configure(state="disabled")
                messagebox.showwarning("Không có dữ liệu", "Không tìm thấy lịch dạy của giáo viên này.")
                return

            self.export_button.configure(state="normal")
            self.status_var.set(
                f"Giáo viên {teacher_label}: {len(self.filtered_records)} tiết, đã sắp xếp "
                f"Thứ → Sáng/Chiều → Tiết; trùng tiết: "
                f"{self.enrichment_result.duplicate_slots}; thiếu PPCT: "
                f"{self.enrichment_result.missing_ppct}; thiếu tên bài: "
                f"{self.enrichment_result.missing_lesson}; dòng trùng do khác tên đã gộp: "
                f"{alias_duplicates_removed}."
            )
            schedule_alerts = collect_schedule_alerts(self.filtered_records)
            schedule_alerts.extend(
                collect_alias_excess_alerts(self.merged_result.records, teachers)
            )
            schedule_alerts.sort(key=lambda alert: alert.record.sort_key)
            if show_dialog:
                self._show_schedule_alert_table(schedule_alerts, teacher_label)
            warnings: list[str] = []
            if self.enrichment_result.missing_lesson:
                warnings.append(
                    f"Thiếu tên bài ở {self.enrichment_result.missing_lesson} dòng."
                )
            if warnings and show_dialog and not schedule_alerts:
                messagebox.showwarning(
                    "Cảnh báo dữ liệu",
                    "\n".join(warnings)
                    + "\n\nCác dòng tương ứng đã được đánh dấu tại cột Ghi chú.",
                )
        except Exception as exc:
            self.enrichment_result = None
            self.export_button.configure(state="disabled")
            messagebox.showerror("Không thể ghép dữ liệu", str(exc))

    def _export_all_teacher_conflicts(self) -> None:
        if self.merged_result is None:
            return
        alerts = collect_all_teacher_duplicate_alerts(self.merged_result.records)
        if not alerts:
            messagebox.showinfo(
                "Không có tiết trùng",
                "Không phát hiện giáo viên nào dạy nhiều lớp trong cùng một khung giờ.",
            )
            return

        output_path = filedialog.asksaveasfilename(
            title="Lưu danh sách tất cả tiết trùng",
            defaultextension=".xlsx",
            initialfile="Danh_sach_tiet_trung_tat_ca_GV.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not output_path:
            return

        try:
            monday = parse_date(self.monday_var.get())
            export_all_teacher_conflicts(alerts, output_path, monday)
            duplicate_slots = len(
                {
                    (
                        alert.record.teacher,
                        alert.record.day_index,
                        alert.record.session,
                        alert.record.period,
                    )
                    for alert in alerts
                }
            )
            duplicate_teachers = len({alert.record.teacher for alert in alerts})
            messagebox.showinfo(
                "Đã xuất báo cáo",
                f"Đã xuất {duplicate_slots} khung tiết trùng của "
                f"{duplicate_teachers} giáo viên.",
            )
        except Exception as exc:
            messagebox.showerror("Không thể xuất báo cáo tiết trùng", str(exc))

    def _export(self) -> None:
        if self.enrichment_result is None or not self.enrichment_result.rows:
            return
        if not self.template_path.get():
            messagebox.showerror("Thiếu tệp A4", "Vui lòng chọn tệp mẫu A4.")
            return
        try:
            monday = parse_date(self.monday_var.get())
            teachers = self.selected_teachers or [self.teacher_var.get()]
            teachers = [compact_text(name) for name in teachers if compact_text(name)]
            week = compact_text(self.week_var.get())
            if len(teachers) == 1:
                teacher_file_label = safe_filename(teachers[0])
            else:
                teacher_file_label = safe_filename("_".join(teachers[:3]))
                if len(teachers) > 3:
                    teacher_file_label += f"_va_{len(teachers) - 3}_ten"
            default_name = f"LBG_{teacher_file_label}"
            if week:
                default_name += f"_Tuan_{week}"
            default_name += ".xlsx"

            output_path = filedialog.asksaveasfilename(
                title="Lưu lịch báo giảng",
                defaultextension=".xlsx",
                initialfile=default_name,
                filetypes=[("Excel", "*.xlsx")],
            )
            if not output_path:
                return

            export_lbg(
                rows=self.enrichment_result.rows,
                output_path=output_path,
                template_file=self.template_path.get(),
                week=week,
                monday=monday,
            )
            self.status_var.set(f"Đã xuất: {output_path}")
            messagebox.showinfo("Hoàn thành", "Đã xuất lịch báo giảng thành công.")
        except Exception as exc:
            messagebox.showerror("Không thể xuất tệp", str(exc))


def run_check(file_paths: list[str], teacher: str | None) -> int:
    results = [parse_schedule_file(path) for path in file_paths]
    merged = merge_parse_results(results)
    teacher_names = (
        [name for name in re.split(r"[,;]", teacher) if compact_text(name)]
        if teacher
        else []
    )
    teacher_names = [compact_text(name) for name in teacher_names]
    selected = filter_teachers(merged.records, teacher_names)
    report = {
        "files": file_paths,
        "sections": merged.section_count,
        "records": len(merged.records),
        "duplicate_records_removed": merged.duplicate_records,
        "teachers": len(merged.teachers),
        "ignored_hidden_rows": merged.ignored_hidden_rows,
        "ignored_hidden_columns": merged.ignored_hidden_columns,
        "missing_period_rows": merged.missing_period_rows,
        "week": merged.metadata.week,
        "monday": merged.metadata.monday.isoformat() if merged.metadata.monday else None,
        "teacher": teacher,
        "selected_teacher_names": teacher_names,
        "teacher_records": [
            {
                "day": record.day_name,
                "session": record.session,
                "period": record.period,
                "class": record.class_name,
                "subject": record.subject,
                "teacher": record.teacher,
                "school": record.school_code or "Trường chính",
                "period_inferred": record.period_inferred,
                "source_row": record.source_row,
                "source_file": record.source_file,
            }
            for record in selected
        ],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Lọc nhiều tệp thời khóa biểu và xuất lịch báo giảng."
    )
    parser.add_argument(
        "--check",
        nargs="+",
        metavar="FILE_TKB",
        help="Kiểm tra bộ đọc không mở giao diện.",
    )
    parser.add_argument("--teacher", help="Tên giáo viên cần in khi dùng --check.")
    args = parser.parse_args()

    if args.check:
        return run_check(args.check, args.teacher)

    app = TimetableFilterApp()
    app.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
